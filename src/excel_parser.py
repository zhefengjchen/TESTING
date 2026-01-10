"""
Excel Parser module for Invoice Management System.
Handles parsing of Excel files from testing labs.
"""

import os
import re
from typing import List, Dict, Any, Tuple, Optional
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


class ExcelParser:
    """Parser for Excel invoice files from testing labs."""

    # Excel header to database column mapping
    HEADER_MAPPING = {
        "Tajan Bidding Tracking Number": "tajan_bidding_tracking_number",
        "Amazon Test Request #": "amazon_test_request",
        "Amazon Tracker Number": "amazon_tracker_number",
        "ASIN#": "asin",
        "Development Center (SEA, LUX/EU, JP)": "development_center",
        "Category (New HCC)": "category",
        "Product Brand": "product_brand",
        "Product Description": "product_description",
        "Testing SLA (Working days)": "testing_sla",
        "Test / Service Type": "test_service_type",
        "Quotation/Order Number": "quotation_order_number",
        "Request Date": "request_date",
        "Test Start Date": "test_start_date",
        "Report Delivered Date": "report_delivered_date",
        "Report number": "report_number",
        "Test / inspection Location": "test_inspection_location",
        "Product line": "product_line",
        "AMAZON QUALITY MANAGER": "amazon_quality_manager",
        "AMAZON SOURCING MANAGER": "amazon_sourcing_manager",
        "Invoice #": "invoice_number",
        "Lab contact": "lab_contact",
        "COMMENT (IF ANY)": "comment",
        "AMOUNT (Currency=USD)": "amount_usd",
    }

    INSPECTION_HEADER_MAPPING = {
        "Inspection ID": "inspection_id",
        "Factory ID": "factory_id",
        "Amazon Tracker Number": "amazon_tracker_number",
        "Factory Name": "factory_name",
        "Product Description": "product_description",
        "Manday": "manday",
        "Test / Service Type": "test_service_type",
        "Quotation/Order Number": "quotation_order_number",
        "Request Date": "request_date",
        "Test Start Date": "test_start_date",
        "Report Delivered Date": "report_delivered_date",
        "Report number": "report_number",
        "Test / inspection Location": "test_inspection_location",
        "Product line": "product_line",
        "AMAZON QUALITY MANAGER": "amazon_quality_manager",
        "AMAZON SOURCING MANAGER": "amazon_sourcing_manager",
        "Invoice #": "invoice_number",
        "Lab contact": "lab_contact",
        "COMMENT (IF ANY)": "comment",
        "AMOUNT (Currency=USD)": "amount_usd",
    }

    # Rows to skip (subtotals, tax, etc.) - case insensitive regex patterns
    SKIP_PATTERNS = [
        r'^remit\s*payment\s*to',
        r'^subtotal',
        r'^tax\s*rate',
        r'^sales\s*tax',
        r'^freight',
        r'^total$',
        r'^grand\s*total',
    ]

    # Exact values to skip (case insensitive)
    SKIP_VALUES = [
        "REMIT PAYMENT TO:",
        "SUBTOTAL",
        "TAX RATE",
        "SALES TAX",
        "FREIGHT",
        "TOTAL",
    ]

    # Sheet name to parse
    TARGET_SHEET = "TESTING+SERVICES"
    INSPECTION_SHEET = "INSPECTION"

    def __init__(self, lab_names: List[str] = None):
        """Initialize the parser with known lab names."""
        self.lab_names = lab_names or ["BV", "ITS", "TUV", "SGS"]

    def detect_lab_name(self, filename: str) -> Optional[str]:
        """
        Detect lab name from filename.
        Returns the lab name if found, None otherwise.
        """
        filename_upper = filename.upper()
        for lab in self.lab_names:
            if lab.upper() in filename_upper:
                return lab.upper()
        return None

    def parse_file(self, file_path: str, invoice_date: str = None) -> Tuple[List[Dict[str, Any]], Optional[str], Optional[str]]:
        """
        Parse an Excel file and extract invoice data.

        Args:
            file_path: Path to the Excel file
            invoice_date: Optional invoice date to add to all records

        Returns:
            Tuple of (list of invoice records, detected lab name, error message if any)
        """
        filename = os.path.basename(file_path)
        lab_name = self.detect_lab_name(filename)

        try:
            workbook = load_workbook(file_path, data_only=True)
        except Exception as e:
            return [], lab_name, f"Failed to open Excel file: {str(e)}"

        # Find the TESTING+SERVICES sheet
        sheet = None
        for sheet_name in workbook.sheetnames:
            if sheet_name.upper() == self.TARGET_SHEET.upper():
                sheet = workbook[sheet_name]
                break

        if sheet is None:
            # Try to find a sheet containing the target name
            for sheet_name in workbook.sheetnames:
                if self.TARGET_SHEET.upper() in sheet_name.upper():
                    sheet = workbook[sheet_name]
                    break

        if sheet is None:
            return [], lab_name, f"Sheet '{self.TARGET_SHEET}' not found. Available sheets: {', '.join(workbook.sheetnames)}"

        # Parse the sheet
        records = self._parse_sheet(sheet, lab_name, invoice_date)

        workbook.close()
        return records, lab_name, None

    def parse_inspection_file(
        self,
        file_path: str,
        invoice_date: str = None
    ) -> Tuple[List[Dict[str, Any]], Optional[str], Optional[str]]:
        """
        Parse an Excel file and extract inspection data.

        Args:
            file_path: Path to the Excel file
            invoice_date: Optional invoice date to add to all records

        Returns:
            Tuple of (list of inspection records, detected lab name, error message if any)
        """
        filename = os.path.basename(file_path)
        lab_name = self.detect_lab_name(filename)

        try:
            workbook = load_workbook(file_path, data_only=True)
        except Exception as e:
            return [], lab_name, f"Failed to open Excel file: {str(e)}"

        sheet = None
        for sheet_name in workbook.sheetnames:
            if sheet_name.upper() == self.INSPECTION_SHEET.upper():
                sheet = workbook[sheet_name]
                break

        if sheet is None:
            for sheet_name in workbook.sheetnames:
                if self.INSPECTION_SHEET.upper() in sheet_name.upper():
                    sheet = workbook[sheet_name]
                    break

        if sheet is None:
            workbook.close()
            return [], lab_name, None

        records = self._parse_inspection_sheet(sheet, lab_name, invoice_date)
        workbook.close()
        return records, lab_name, None

    def _parse_sheet(self, sheet: Worksheet, lab_name: Optional[str], invoice_date: Optional[str]) -> List[Dict[str, Any]]:
        """Parse a worksheet and extract invoice records."""
        records = []

        # Find header row
        header_row_idx = None
        header_mapping = {}

        normalized_header_mapping = self._normalized_header_mapping(self.HEADER_MAPPING)

        for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=50), start=1):
            row_values = [cell.value for cell in row]

            # Check if this row contains expected headers
            matching_headers = 0
            for col_idx, cell_value in enumerate(row_values):
                if cell_value and isinstance(cell_value, str):
                    cell_value_stripped = self._normalize_header(cell_value)
                    if cell_value_stripped in normalized_header_mapping:
                        matching_headers += 1
                        header_mapping[col_idx] = normalized_header_mapping[cell_value_stripped]

            # If we found enough headers, this is our header row
            if matching_headers >= 5:  # At least 5 matching headers
                header_row_idx = row_idx
                break

        if header_row_idx is None:
            return records

        # Parse data rows
        for row in sheet.iter_rows(min_row=header_row_idx + 1):
            row_values = [cell.value for cell in row]

            # Skip if row is completely empty
            if all(v is None or (isinstance(v, str) and v.strip() == '') for v in row_values):
                continue

            # Check if this is a subtotal/tax/total row
            first_non_empty = None
            for v in row_values:
                if v is not None and (not isinstance(v, str) or v.strip() != ''):
                    first_non_empty = str(v).strip().lower()
                    break

            if first_non_empty and self._should_skip_row(first_non_empty):
                continue

            # Build record from row
            record = {}
            for col_idx, db_column in header_mapping.items():
                if col_idx < len(row_values):
                    value = row_values[col_idx]
                    record[db_column] = self._format_value(value, db_column)

            # Add lab name and invoice date
            record['lab_name'] = lab_name or ''
            record['invoice_date'] = invoice_date or ''

            # Skip if the record is essentially empty (no meaningful data)
            if not self._has_meaningful_data(record):
                continue

            records.append(record)

        return records

    def _parse_inspection_sheet(
        self,
        sheet: Worksheet,
        lab_name: Optional[str],
        invoice_date: Optional[str]
    ) -> List[Dict[str, Any]]:
        """Parse a worksheet and extract inspection records."""
        records = []
        header_row_idx = None
        header_mapping = {}
        normalized_header_mapping = self._normalized_header_mapping(self.INSPECTION_HEADER_MAPPING)

        for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=50), start=1):
            row_values = [cell.value for cell in row]
            matching_headers = 0
            for col_idx, cell_value in enumerate(row_values):
                if cell_value and isinstance(cell_value, str):
                    cell_value_stripped = self._normalize_header(cell_value)
                    if cell_value_stripped in normalized_header_mapping:
                        matching_headers += 1
                        header_mapping[col_idx] = normalized_header_mapping[cell_value_stripped]

            if matching_headers >= 5:
                header_row_idx = row_idx
                break

        if header_row_idx is None:
            return records

        for row in sheet.iter_rows(min_row=header_row_idx + 1):
            row_values = [cell.value for cell in row]
            if all(v is None or (isinstance(v, str) and v.strip() == '') for v in row_values):
                continue

            first_non_empty = None
            for v in row_values:
                if v is not None and (not isinstance(v, str) or v.strip() != ''):
                    first_non_empty = str(v).strip().lower()
                    break

            if first_non_empty and self._should_skip_row(first_non_empty):
                continue

            record = {}
            for col_idx, db_column in header_mapping.items():
                if col_idx < len(row_values):
                    value = row_values[col_idx]
                    record[db_column] = self._format_value(value, db_column)

            record['test_service_type'] = 'PSI'
            record['lab_name'] = lab_name or ''
            record['invoice_date'] = invoice_date or ''

            amount_value = record.get('amount_usd', 0)
            if not amount_value or float(amount_value) <= 0:
                continue

            raw_amount = record.get('amount_usd', '')
            if isinstance(raw_amount, str) and self._should_skip_row(raw_amount.strip().lower()):
                continue

            if not self._has_meaningful_inspection_data(record):
                continue

            records.append(record)

        return records

    def _should_skip_row(self, first_value: str) -> bool:
        """Check if a row should be skipped based on its first value."""
        # Check exact values (case insensitive)
        first_value_upper = first_value.strip().upper()
        for skip_value in self.SKIP_VALUES:
            if first_value_upper == skip_value.upper():
                return True

        # Check regex patterns (case insensitive)
        for pattern in self.SKIP_PATTERNS:
            if re.match(pattern, first_value, re.IGNORECASE):
                return True
        return False

    @staticmethod
    def _normalize_header(value: Any) -> str:
        """Normalize header cell values for matching."""
        if value is None:
            return ""
        return re.sub(r"\s+", " ", str(value)).strip()

    def _normalized_header_mapping(self, mapping: Dict[str, str]) -> Dict[str, str]:
        """Return a mapping with normalized keys for header matching."""
        return {self._normalize_header(key): value for key, value in mapping.items()}

    def _format_value(self, value: Any, column_name: str) -> Any:
        """Format a cell value based on the column type."""
        if value is None:
            return ''

        # Handle date columns
        if column_name in ('request_date', 'test_start_date', 'report_delivered_date', 'invoice_date'):
            if hasattr(value, 'strftime'):
                return value.strftime('%Y-%m-%d')
            return str(value).strip() if value else ''

        # Handle amount column
        if column_name == 'amount_usd':
            if isinstance(value, (int, float)):
                return float(value)
            if isinstance(value, str):
                # Try to parse string as number
                try:
                    cleaned = re.sub(r'[^\d.-]', '', value)
                    return float(cleaned) if cleaned else 0.0
                except ValueError:
                    return 0.0
            return 0.0

        # Default: convert to string
        return str(value).strip() if value else ''

    def _has_meaningful_data(self, record: Dict[str, Any]) -> bool:
        """Check if a record has meaningful data (not just empty values)."""
        # Check for at least one non-empty key field
        key_fields = [
            'tajan_bidding_tracking_number',
            'amazon_test_request',
            'asin',
            'invoice_number',
            'test_service_type'
        ]

        for field in key_fields:
            value = record.get(field, '')
            if value and str(value).strip():
                return True

        # Also check if there's an amount
        amount = record.get('amount_usd', 0)
        if amount and float(amount) > 0:
            return True

        return False

    def _has_meaningful_inspection_data(self, record: Dict[str, Any]) -> bool:
        """Check if an inspection record has meaningful data."""
        key_fields = [
            'inspection_id',
            'factory_id',
            'amazon_tracker_number',
            'invoice_number',
        ]

        for field in key_fields:
            value = record.get(field, '')
            if value and str(value).strip():
                return True

        amount = record.get('amount_usd', 0)
        if amount and float(amount) > 0:
            return True

        return False


class DataValidator:
    """Validates invoice data according to business rules."""

    def __init__(self, allowed_test_types: List[str]):
        """Initialize validator with allowed test types."""
        self.allowed_test_types = {
            self._normalize_test_type(t) for t in allowed_test_types if t is not None
        }

    @staticmethod
    def _normalize_test_type(value: str) -> str:
        """Normalize test type values for comparison."""
        if value is None:
            return ""
        normalized = re.sub(r"\s+", " ", str(value)).strip().lower()
        return normalized

    def validate_record(self, record: Dict[str, Any]) -> Tuple[bool, List[str]]:
        """
        Validate a single invoice record.

        Returns:
            Tuple of (is_valid, list of validation errors)
        """
        errors = []

        # Check for missing tracking number
        tracking_number = record.get('tajan_bidding_tracking_number', '')
        if not tracking_number or str(tracking_number).strip() == '':
            errors.append("Missing Tracking Number")

        # Check for invalid test type
        test_type = record.get('test_service_type', '')
        if test_type:
            test_type_normalized = self._normalize_test_type(test_type)
            if test_type_normalized and test_type_normalized not in self.allowed_test_types:
                errors.append(f"Invalid Test Type: {test_type}")
        else:
            # Empty test type is also invalid
            errors.append("Missing Test/Service Type")

        is_valid = len(errors) == 0
        return is_valid, errors

    def validate_records(self, records: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], List[Tuple[Dict[str, Any], str]]]:
        """
        Validate a list of invoice records.

        Returns:
            Tuple of (valid_records, list of (invalid_record, error_message) tuples)
        """
        valid_records = []
        invalid_records = []

        for record in records:
            is_valid, errors = self.validate_record(record)
            if is_valid:
                valid_records.append(record)
            else:
                error_message = "; ".join(errors)
                invalid_records.append((record, error_message))

        return valid_records, invalid_records
